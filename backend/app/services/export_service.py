import io
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from fastapi import HTTPException
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


class ExportService:
    """Build download files in memory; drafts never need to be stored on disk."""

    _pdf_font_name = "NirmalaUI"
    _font_registered = False

    @classmethod
    def _register_pdf_font(cls) -> str:
        if cls._font_registered:
            return cls._pdf_font_name
        for candidate in (
            Path(r"C:\Windows\Fonts\Nirmala.ttc"),
            Path(r"C:\Windows\Fonts\Nirmala.ttf"),
        ):
            if candidate.is_file():
                try:
                    pdfmetrics.registerFont(TTFont(cls._pdf_font_name, str(candidate)))
                    cls._font_registered = True
                    return cls._pdf_font_name
                except Exception:
                    continue
        return "Helvetica"

    @staticmethod
    def _text_bytes(text: str) -> bytes:
        # UTF-8 BOM lets Windows Notepad open Hindi and English text reliably.
        return text.encode("utf-8-sig")

    @staticmethod
    def _docx_bytes(text: str) -> bytes:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        normal = document.styles["Normal"]
        normal.font.name = "Nirmala UI"
        normal.font.size = Pt(10.5)
        for line in text.splitlines() or [""]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(line)
            run.font.name = "Nirmala UI"
            run.font.size = Pt(10.5)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @classmethod
    def _pdf_bytes(cls, text: str) -> bytes:
        buffer = io.BytesIO()
        font_name = cls._register_pdf_font()
        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=0.65 * inch,
            rightMargin=0.65 * inch,
            topMargin=0.65 * inch,
            bottomMargin=0.65 * inch,
        )
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "DraftBody",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=10,
            leading=15,
            alignment=TA_LEFT,
            spaceAfter=6,
        )
        story = []
        for paragraph in text.split("\n\n") or [""]:
            if paragraph.strip():
                safe_text = escape(paragraph).replace("\n", "<br/>")
                story.append(Paragraph(safe_text, body))
            else:
                story.append(Spacer(1, 6))
        document.build(story)
        return buffer.getvalue()

    def build(self, text: str, file_format: str) -> tuple[bytes, str, str]:
        if file_format == "txt":
            return self._text_bytes(text), "text/plain; charset=utf-8", "txt"
        if file_format == "docx":
            return (
                self._docx_bytes(text),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "docx",
            )
        if file_format == "pdf":
            return self._pdf_bytes(text), "application/pdf", "pdf"
        raise HTTPException(422, "Choose TXT, Word (.docx), or PDF.")
